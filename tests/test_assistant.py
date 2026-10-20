"""Tests for parsing, matching, and the chat turns a reviewer will actually try."""

from __future__ import annotations

import json
import os
from datetime import datetime, timezone
from io import BytesIO
from pathlib import Path

import pytest
from docx import Document
from fastapi.testclient import TestClient

from app.agents.resume import parse_resume_text
from app.ingest import normalize_ashby
from app.main import create_app
from app.skills import extract_skills
from app.textutil import experience_required, plain_text, years_between
from scripts.make_sample_resume import LINES, build_pdf

ROOT = Path(__file__).resolve().parents[1]
SAMPLE = "\n".join(LINES)
# Tests stay on the phrase rules even if a developer key is loaded.
os.environ.pop("GEMINI_API_KEY", None)


def test_javascript_is_not_java():
    assert extract_skills("JavaScript developer") == ["javascript"]
    both = extract_skills("I write Java and JavaScript")
    assert "java" in both and "javascript" in both
    assert extract_skills("json file") == []


def test_spaced_stack_is_javascript_typescript_and_node(client):
    api, _data = client
    started = api.post(
        "/api/chat",
        data={"message": "show only remote roles"},
        files={"file": ("alex.docx", _docx_bytes(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
    ).json()
    asked = api.post(
        "/api/chat",
        data={"message": "js type script node.js", "session_id": started["session_id"]},
    ).json()
    assert asked["filters"]["skills"] == ["javascript", "typescript", "node.js"]
    assert asked["filters"]["skill_match"] == "any"
    assert asked["filters"]["remote"] is None
    assert asked["jobs"]
    for job in asked["jobs"]:
        skills = set(job["skills"])
        assert skills & {"javascript", "typescript", "node.js"}
    assert "narrow the roles" not in asked["reply"].lower()


def test_react_native_keeps_react():
    found = set(extract_skills("Built a React Native app"))
    assert "react native" in found
    assert "react" in found


def test_greenhouse_html_is_unescaped():
    text = plain_text("&lt;p&gt;We use Python and PostgreSQL.&lt;/p&gt;")
    assert "Python" in text
    assert "<" not in text
    assert "lt;" not in text


def test_experience_label_skips_age():
    label, years = experience_required("Analyst", "Must be 18 years of age. 3+ years of experience required.")
    assert years == 3
    assert label == "3+ years"


def test_overlapping_jobs_are_not_double_counted():
    # Two roles in the same window should count once.
    today = datetime(2026, 10, 19, tzinfo=timezone.utc)
    total = years_between(
        [((2024, 1), (2026, 1)), ((2024, 6), (2025, 6))],
        today,
    )
    assert total == 2.0


def test_ashby_location_stays_readable():
    job = normalize_ashby(
        {
            "id": "abc",
            "title": "Software Engineer",
            "location": "New York City, NY",
            "secondaryLocations": [
                {"location": "San Francisco, CA", "address": {"postalAddress": {"addressLocality": "San Francisco"}}}
            ],
            "descriptionPlain": "React and Python. 3+ years of experience.",
            "jobUrl": "https://jobs.ashbyhq.com/example/abc",
            "isRemote": False,
        },
        "example",
        "Example",
    )
    assert job is not None
    assert "San Francisco, CA" in job["location"]
    assert "{" not in job["location"]


def test_sample_resume_fields():
    profile = parse_resume_text(SAMPLE, today=datetime(2026, 10, 19, tzinfo=timezone.utc))
    assert profile["name"] == "Alex Rivera"
    assert profile["email"] == "alex.rivera@example.com"
    assert profile["phone"]
    assert profile["designation"] == "Software Engineer"
    assert "react" in profile["skills"]
    assert "typescript" in profile["skills"]
    assert "python" in profile["skills"]
    assert profile["years"] is not None
    assert 2.2 <= float(profile["years"]) <= 3.6
    assert any("Northwind" in (role.get("company") or "") for role in profile["experience"])
    assert any("State University" in line for line in profile["education"])
    assert any("Trailhead" in name for name in profile["projects"])


def _jobs() -> list[dict]:
    titles = [
        ("Frontend Software Engineer", ["react", "typescript", "javascript"], False, "San Francisco", "2+ years", 2),
        ("Software Engineer", ["react", "node.js", "postgresql", "typescript"], True, "Remote - US", "2+ years", 2),
        ("Backend Software Engineer", ["python", "postgresql", "django"], True, "Remote", "3+ years", 3),
        ("Software Engineer", ["react", "typescript", "aws"], True, "Remote - Canada", "1+ years", 1),
        ("Product Designer", ["figma"], False, "New York", "3+ years", 3),
        ("Senior Staff Engineer", ["go", "kubernetes"], True, "Remote", "8+ years", 8),
        ("Software Engineer", ["react", "python", "node.js"], True, "Remote", "2+ years", 2),
        ("Software Engineer", ["typescript", "react", "graphql"], True, "Remote - Europe", "2+ years", 2),
        ("Software Engineer Intern", ["react", "python"], True, "Remote", "Internship", 0),
    ]
    rows = []
    for index, (title, skills, remote, location, label, minimum) in enumerate(titles, start=1):
        rows.append({
            "id": f"gh_example_{index}",
            "title": title,
            "company": f"Example {index}",
            "location": location,
            "experience_required": label,
            "experience_years_min": minimum,
            "skills": skills,
            "description": f"{title} needs {', '.join(skills)}. Location {location}.",
            "source_link": f"https://example.com/jobs/{index}",
            "source": "Greenhouse",
            "remote": remote,
        })
    return rows


@pytest.fixture
def client(tmp_path: Path):
    data = tmp_path / "data"
    data.mkdir()
    payload = {"generated_at": "2026-10-19T00:00:00+00:00", "jobs": _jobs()}
    (data / "jobs.json").write_text(json.dumps(payload), encoding="utf-8")
    app = create_app(data)
    return TestClient(app), data


def _docx_bytes() -> bytes:
    document = Document()
    for line in LINES:
        document.add_paragraph(line)
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def test_chat_reads_resume_and_recommends(client):
    api, _data = client
    response = api.post(
        "/api/chat",
        data={"message": "", "session_id": ""},
        files={"file": ("alex.docx", _docx_bytes(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
    )
    body = response.json()
    assert response.status_code == 200
    assert "Alex Rivera" in body["reply"]
    assert body["profile"]["designation"] == "Software Engineer"
    assert body["jobs"], body["reply"]
    assert all(job["reason"] for job in body["jobs"])
    # The designer role shouldn't outrank a React engineer for this resume.
    titles = [job["title"] for job in body["jobs"]]
    assert "Product Designer" not in titles


def test_remote_filter_apply_and_duplicate(client):
    api, data = client
    first = api.post(
        "/api/chat",
        data={"message": "show only remote roles"},
        files={"file": ("alex.pdf", build_pdf(LINES), "application/pdf")},
    ).json()
    assert first["jobs"]
    assert all(job["remote"] for job in first["jobs"])
    session_id = first["session_id"]

    applied = api.post("/api/chat", data={"message": "apply to job 1", "session_id": session_id}).json()
    assert "not submit" in applied["reply"].lower()
    assert applied["saved_applications"][0]["status"] == "not_submitted"
    assert applied["saved_applications"][0]["job_id"] == first["jobs"][0]["id"]

    again = api.post("/api/chat", data={"message": "apply to job 1", "session_id": session_id}).json()
    assert "not submitted" in again["reply"].lower()
    stored = json.loads((data / "applications.json").read_text(encoding="utf-8"))
    assert len(stored["applications"]) == 1

    status = api.post("/api/chat", data={"message": "show my applied jobs", "session_id": session_id}).json()
    assert first["jobs"][0]["title"] in status["reply"]
    assert "not submitted" in status["reply"].lower() or "none of these were submitted" in status["reply"].lower()


def test_frontend_request_drops_the_remote_filter(client):
    api, _data = client
    started = api.post(
        "/api/chat",
        data={"message": "show only remote roles"},
        files={"file": ("alex.docx", _docx_bytes(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
    ).json()
    assert started["filters"]["remote"] is True
    narrowed = api.post(
        "/api/chat",
        data={"message": "only frontend roles", "session_id": started["session_id"]},
    ).json()
    assert narrowed["filters"]["remote"] is None
    assert narrowed["filters"]["track"] == "frontend"
    assert narrowed["jobs"]
    titles = [job["title"].lower() for job in narrowed["jobs"]]
    assert titles
    assert all("frontend" in title or "front-end" in title or "front end" in title for title in titles)
    assert "remote only" not in narrowed["reply"].lower()


def test_suggest_any_changes_gives_resume_edits(client):
    api, _data = client
    started = api.post(
        "/api/chat",
        data={"message": ""},
        files={"file": ("alex.docx", _docx_bytes(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
    ).json()
    advice = api.post(
        "/api/chat",
        data={"message": "can you suggest any changes?", "session_id": started["session_id"]},
    ).json()
    assert "change the resume" in advice["reply"].lower()
    assert started["jobs"][0]["title"] in advice["reply"]


def test_role_title_searches_matching_jobs(tmp_path: Path):
    data = tmp_path / "data"
    data.mkdir()
    payload = {
        "generated_at": "2026-10-19T00:00:00+00:00",
        "jobs": [
            {
                "id": "gh_example_ai",
                "title": "AI Engineer",
                "company": "Northwind",
                "location": "New York",
                "experience_required": "3+ years",
                "experience_years_min": 3,
                "skills": ["python", "pytorch"],
                "description": "Build models.",
                "source_link": "https://example.com/jobs/ai",
                "source": "Greenhouse",
                "remote": False,
            },
            {
                "id": "gh_example_saas",
                "title": "Senior Software Engineer (Python / AI Insurance SaaS)",
                "company": "Contoso",
                "location": "Remote",
                "experience_required": "5+ years",
                "experience_years_min": 5,
                "skills": ["python"],
                "description": "Insurance platform.",
                "source_link": "https://example.com/jobs/saas",
                "source": "Greenhouse",
                "remote": True,
            },
            {
                "id": "gh_example_ml",
                "title": "Sr. Machine Learning Engineer",
                "company": "Fabrikam",
                "location": "Remote",
                "experience_required": "4+ years",
                "experience_years_min": 4,
                "skills": ["python"],
                "description": "Training pipelines.",
                "source_link": "https://example.com/jobs/ml",
                "source": "Greenhouse",
                "remote": True,
            },
        ],
    }
    (data / "jobs.json").write_text(json.dumps(payload), encoding="utf-8")
    api = TestClient(create_app(data))
    started = api.post(
        "/api/chat",
        data={"message": "show only remote roles"},
        files={"file": ("alex.docx", _docx_bytes(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
    ).json()
    asked = api.post(
        "/api/chat",
        data={"message": "ai engineer", "session_id": started["session_id"]},
    ).json()
    assert asked["filters"]["query"] == "ai engineer"
    assert asked["filters"]["remote"] is None
    titles = [job["title"] for job in asked["jobs"]]
    assert titles
    assert "AI Engineer" in titles
    assert "Sr. Machine Learning Engineer" in titles
    assert all("Insurance" not in title for title in titles)
    assert "narrow the roles" not in asked["reply"].lower()


def test_gemini_reads_the_chat_and_writes_the_reply(tmp_path, monkeypatch):
    monkeypatch.setenv("GEMINI_API_KEY", "test-key")
    data = tmp_path / "data"
    data.mkdir()
    payload = {
        "generated_at": "2026-10-19T00:00:00+00:00",
        "jobs": [
            {
                "id": "gh_example_nw",
                "title": "AI Engineer",
                "company": "Northwind",
                "location": "New York",
                "experience_required": "3+ years",
                "experience_years_min": 3,
                "skills": ["python"],
                "description": "Models.",
                "source_link": "https://example.com/jobs/nw",
                "source": "Greenhouse",
                "remote": False,
            },
            {
                "id": "gh_example_co",
                "title": "Software Engineer",
                "company": "Contoso",
                "location": "Remote",
                "experience_required": "2+ years",
                "experience_years_min": 2,
                "skills": ["react"],
                "description": "Product.",
                "source_link": "https://example.com/jobs/co",
                "source": "Greenhouse",
                "remote": True,
            },
        ],
    }
    (data / "jobs.json").write_text(json.dumps(payload), encoding="utf-8")

    def fake_urlopen(request, timeout=12):
        prompt = json.loads(request.data.decode("utf-8"))["contents"][0]["parts"][0]["text"]
        if prompt.startswith("Read this job-search chat"):
            text = json.dumps({
                "kind": "jobs",
                "company": "Northwind",
                "query": None,
                "remote": None,
                "skills": [],
                "track": None,
            })
        else:
            text = "Northwind is the employer you asked about. That application was not submitted."

        class _Body:
            def read(self):
                return json.dumps({"candidates": [{"content": {"parts": [{"text": text}]}}]}).encode("utf-8")

            def __enter__(self):
                return self

            def __exit__(self, *args):
                return False

        return _Body()

    monkeypatch.setattr("app.agents.llm.urllib.request.urlopen", fake_urlopen)
    api = TestClient(create_app(data))
    started = api.post(
        "/api/chat",
        data={"message": ""},
        files={"file": ("alex.docx", _docx_bytes(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
    ).json()
    asked = api.post(
        "/api/chat",
        data={"message": "what about that company", "session_id": started["session_id"]},
    ).json()
    assert asked["filters"]["company"] == "Northwind"
    assert asked["jobs"]
    assert all(job["company"] == "Northwind" for job in asked["jobs"])
    assert "Northwind" in asked["reply"]
    assert "Read the message" in asked["activity"]


def test_reply_refuses_a_false_submission():
    from app.agents.llm import _accept_reply

    fallback = "I did not submit the form."
    assert _accept_reply("I submitted your application.", fallback, {"must_not_claim_submitted": True}) == fallback
    assert _accept_reply("The form was not submitted.", fallback, {"must_not_claim_submitted": True}) == "The form was not submitted."


def test_invalid_job_number(client):
    api, _data = client
    started = api.post(
        "/api/chat",
        data={"message": ""},
        files={"file": ("alex.docx", _docx_bytes(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
    ).json()
    missed = api.post(
        "/api/chat",
        data={"message": "apply to job 9", "session_id": started["session_id"]},
    ).json()
    assert "isn't on the list" in missed["reply"]
    assert missed["saved_applications"] == []


def _letter_docx() -> bytes:
    document = Document()
    document.add_paragraph("Quarterly facilities note")
    document.add_paragraph(
        "The warehouse move finished on Friday. Please review the invoice totals and send comments to the operations desk before Monday."
    )
    document.add_paragraph("Parking passes for the new building will be handed out next week by the facilities team.")
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def test_unreadable_resume(client):
    api, _data = client
    response = api.post(
        "/api/chat",
        data={"message": ""},
        files={"file": ("scan.pdf", b"%PDF-1.4\n1 0 obj\n<<>>\nendobj\ntrailer\n<<>>\n%%EOF", "application/pdf")},
    )
    assert response.status_code == 200
    reply = response.json()["reply"].lower()
    assert "couldn't" in reply or "text" in reply


def test_other_document_is_not_treated_as_a_resume(client):
    api, _data = client
    uploaded = api.post(
        "/api/chat",
        data={"message": ""},
        files={"file": ("notes.docx", _letter_docx(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
    ).json()
    assert uploaded["profile"] is None
    assert uploaded["jobs"] == []
    assert "doesn't look like a resume" in uploaded["reply"].lower()
    asked = api.post(
        "/api/chat",
        data={"message": "show remote jobs", "session_id": uploaded["session_id"]},
    ).json()
    assert asked["jobs"] == []
    assert asked["profile"] is None
    assert "won't" in asked["reply"].lower() or "not going to score" in asked["reply"].lower()


def test_no_skill_overlap_filter(client):
    api, _data = client
    started = api.post(
        "/api/chat",
        data={"message": ""},
        files={"file": ("alex.docx", _docx_bytes(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
    ).json()
    narrowed = api.post(
        "/api/chat",
        data={"message": "jobs needing Figma", "session_id": started["session_id"]},
    ).json()
    # Figma is on the designer role only. It can show up as a stretch, or as an empty set
    # if the filter wipes the list. Either way the reply has to say what happened.
    if narrowed["jobs"]:
        assert all("figma" in " ".join(job["skills"]) for job in narrowed["jobs"])
    else:
        assert "nothing" in narrowed["reply"].lower() or "figma" in narrowed["reply"].lower()


def test_shipped_dataset():
    path = ROOT / "data" / "jobs.json"
    if not path.exists():
        pytest.skip("dataset not built yet")
    raw = json.loads(path.read_text(encoding="utf-8"))
    jobs = raw["jobs"]
    assert 1500 <= len(jobs) <= 2000
    ids = set()
    required = {
        "id", "title", "company", "location", "experience_required",
        "skills", "description", "source_link",
    }
    for job in jobs:
        assert required <= set(job)
        assert job["id"] not in ids
        ids.add(job["id"])
        link = job["source_link"]
        assert link.startswith("http")
        host = link.split("/")[2].lower()
        assert "example.com" not in host
        job_id = job["id"]
        if job_id.startswith("gh_"):
            assert "greenhouse.io" in host or "gh_jid=" in link, link
        elif job_id.startswith("lv_"):
            assert "lever.co" in host, link
        elif job_id.startswith("ab_"):
            assert "ashbyhq.com" in host, link
        else:
            raise AssertionError(job_id)
        assert isinstance(job["skills"], list)
        assert job["title"] and job["company"]
